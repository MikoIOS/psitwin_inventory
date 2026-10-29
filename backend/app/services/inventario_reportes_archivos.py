"""In-memory report renderers. Both download and mail use this implementation."""
from io import BytesIO
from datetime import datetime, timezone
from html import escape
from zoneinfo import ZoneInfo

COLUMNS = [('codigo', 'Código'), ('producto', 'Producto'), ('familia', 'Familia'), ('categoria', 'Categoría'),
    ('unidad', 'Unidad'), ('stock_actual', 'Stock actual'), ('stock_minimo', 'Mínimo'), ('stock_maximo', 'Máximo'),
    ('estado', 'Estado'), ('almacen', 'Almacén'), ('zona', 'Zona'), ('rack', 'Rack'), ('nivel', 'Nivel'), ('ultima_actualizacion', 'Actualizado')]


def local_date(value):
    return (value if value.tzinfo else value.replace(tzinfo=timezone.utc)).astimezone(ZoneInfo('America/Lima')).replace(tzinfo=None) if isinstance(value, datetime) else value

def label(value):
    if value is None: return 'Sin configurar'
    if isinstance(value, datetime): return local_date(value).strftime('%d/%m/%Y %H:%M')
    if isinstance(value, list): return ', '.join(str(v) for v in value)
    return str(value)


def excel(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    wb = Workbook(); wb.remove(wb.active)
    def sheet(name, columns, records):
        ws = wb.create_sheet(name)
        ws.append(['PsiTwin Inventory - Resumen de inventario'])
        ws.append(['Empresa', data['empresa'], 'Alcance', data['alcance']])
        ws.append(['Generado', local_date(data['generado_en']), 'Solicitado por', data['solicitante']])
        ws.append(['Origen: existencias y alertas registradas al generar el reporte.'])
        ws.append([]); ws.append([title for _, title in columns])
        for row in records:
            values = []
            for key, _ in columns:
                value = local_date(row.get(key))
                if isinstance(value, list): value = ', '.join(map(str, value))
                values.append(value)
            ws.append(values)
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub('', cell.value)
                    cell.data_type = 's'  # Never execute product names/comments as Excel formulas.
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(name='Calibri', size=11, color='14263D')
                if isinstance(cell.value, float): cell.number_format = '#,##0.###'
                if isinstance(cell.value, datetime): cell.number_format = 'dd/mm/yyyy hh:mm'
        for row_number in (1, 6):
            for cell in ws[row_number]:
                cell.fill = PatternFill('solid', fgColor='14263D')
                cell.font = Font(name='Calibri', size=14 if row_number == 1 else 11, bold=True, color='FFFFFF')
            ws.row_dimensions[row_number].height = 30
        for i, (key, title) in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = 34 if key in {'producto','observacion','ubicacion'} else 25 if key in {'codigo','almacen','zona','ultima_actualizacion','indicador'} else 19
        ws.column_dimensions['A'].width = max(ws.column_dimensions['A'].width, 32)
        # Persist row heights: desktop Excel does not auto-fit wrapped rows on open.
        import math
        for row in ws.iter_rows(min_row=7):
            lines = max((sum(max(1, math.ceil(len(part) / max(8, ws.column_dimensions[cell.column_letter].width - 2))) for part in str(cell.value or '').split('\n')) for cell in row), default=1)
            ws.row_dimensions[row[0].row].height = max(25, lines * 15 + 8)
        ws.row_dimensions[2].height = 58
        ws.row_dimensions[3].height = 35
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max(4,len(columns)))
        ws.merge_cells(start_row=4,start_column=1,end_row=4,end_column=max(4,len(columns)))
        ws.row_dimensions[4].height = 30
        ws.freeze_panes = 'A7'
        ws.auto_filter.ref = f'A6:{get_column_letter(len(columns))}{max(6, ws.max_row)}'
        ws.sheet_view.showGridLines = False
        ws.print_title_rows = '1:6'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        return ws
    metrics = [{'indicador': key.replace('_', ' ').capitalize(), 'valor': value} for key, value in data['resumen'].items() if key != 'totales_por_unidad']
    metrics.extend({'indicador': 'Stock total (' + unit + ')', 'valor': value} for unit, value in data['resumen']['totales_por_unidad'].items())
    sheet('Resumen', [('indicador','Indicador'), ('valor','Valor')], metrics)
    sheet('Por producto', [(k,v) for k,v in COLUMNS[:8]] + [('estados','Estados por ubicación'), ('ubicaciones','Ubicaciones')], data['productos'])
    sheet('Por ubicación', COLUMNS, data['ubicaciones'])
    sheet('Alertas', [('producto','Producto'), ('tipo','Tipo'), ('cantidad','Cantidad'), ('unidad','Unidad'), ('umbral','Umbral'), ('ubicacion','Ubicación'), ('observacion','Observación')], data['alertas'])
    from app.services.warehouse_map_service import estructura_reporte
    sheet('Plano', [('almacen','Almacen'), ('zona','Zona'), ('rack','Rack'), ('tipo_rack','Tipo de rack'), ('niveles','Niveles'), ('nivel','Nivel'), ('productos','Productos'), ('monitoreo','Monitoreo'), ('estado','Estado de zona / rack / nivel')], estructura_reporte(data))
    out = BytesIO(); wb.save(out); return out.getvalue()


def pdf(data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, KeepTogether
    out = BytesIO(); styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Cell', fontName='Helvetica', fontSize=8, leading=10, spaceAfter=2))
    styles.add(ParagraphStyle(name='SmallNote', fontSize=9, leading=12, textColor=colors.HexColor('#40526A')))
    def p(value): return Paragraph(escape(label(value)).replace('\n', '<br/>'), styles['Cell'])
    def table(headers, rows, widths):
        cells = [[p(h) for h in headers]] + [[p(v) for v in row] for row in rows]
        result = LongTable(cells, colWidths=widths, repeatRows=1, hAlign='LEFT')
        result.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#D9E8F5')), ('VALIGN',(0,0),(-1,-1),'TOP'),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F1F5F9')]), ('LINEBELOW',(0,0),(-1,0),1,colors.HexColor('#1F4F7A')),
            ('LEFTPADDING',(0,0),(-1,-1),6), ('RIGHTPADDING',(0,0),(-1,-1),6), ('TOPPADDING',(0,0),(-1,-1),6), ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        return result
    story = [Paragraph('PsiTwin Inventory', styles['Title']), Paragraph('Resumen de inventario', styles['Heading1']),
        p('Empresa: ' + data['empresa']), p('Alcance: ' + data['alcance']), p('Generado: ' + label(data['generado_en']) + ' - Solicitado por: ' + data['solicitante']), Spacer(1,12)]
    metrics = data['resumen']
    story += [Paragraph('Resumen de existencias', styles['Heading2']),
        p(f"Productos distintos: {metrics['productos_distintos']} | Con stock: {metrics['productos_con_stock']} | Agotados: {metrics['productos_agotados']} | Stock bajo o en mínimo: {metrics['productos_stock_bajo']} | Con alertas: {metrics['productos_con_alertas']} | Ubicaciones utilizadas: {metrics['ubicaciones_utilizadas']}"),
        p('Cantidades por unidad: ' + (', '.join(f'{value:g} {unit}' for unit,value in metrics['totales_por_unidad'].items()) or 'Sin existencias registradas')), Spacer(1,10), Paragraph('Inventario por ubicación', styles['Heading2'])]
    rows = []
    for r in data['ubicaciones']:
        product = '\n'.join(str(r.get(k) or '') for k in ('producto','codigo','familia','categoria'))
        location = ' / '.join(label(r.get(k)) if r.get(k) is not None else 'Sin asignar' for k in ('almacen','zona','rack','nivel'))
        rows.append([product, r['unidad'], r['stock_actual'], r['stock_minimo'], r['stock_maximo'], r['estado'], location, r['ultima_actualizacion']])
    story += [table(['Producto / código / clasificación','Unidad','Stock','Mín.','Máx.','Estado','Almacén / Zona / Rack / Nivel','Actualizado'], rows, [140,48,42,52,52,85,158,95]) if rows else p('El inventario está vacío para los filtros seleccionados.'), Spacer(1,12), Paragraph('Alertas registradas y estados de stock', ParagraphStyle(name='AlertHeading',parent=styles['Heading2'],keepWithNext=True))]
    if data['alertas']:
        story.append(table(['Producto','Tipo','Cantidad / umbral','Ubicación','Observación'], [[a.get('producto'), a.get('tipo'), f"{label(a.get('cantidad'))} {a.get('unidad') or ''} / {label(a.get('umbral'))}", a.get('ubicacion'), a.get('observacion')] for a in data['alertas']], [130,110,110,210,112]))
    else: story.append(p('No hay alertas para este alcance.'))
    story += [Spacer(1,10), Paragraph('El reporte conserva las unidades de origen y no incluye valoraciones monetarias no registradas. Los umbrales se aplican por ubicación; «Sin configurar» no significa cero.', styles['SmallNote'])]
    from app.services.warehouse_map_service import estructura_reporte
    structure = estructura_reporte(data)
    story += [Spacer(1,12), Paragraph('Plano del almacen - distribucion por nivel', styles['Heading2'])]
    story.append(table(['Almacen / Zona', 'Rack / tipo / niveles', 'Nivel', 'Productos', 'Monitoreo'],
        [[r['almacen'] + ' / ' + r['zona'], r['rack'] + ' / ' + r['tipo_rack'] + ' / ' + str(r['niveles']), r['nivel'], r['productos'], r['monitoreo'] + ' / ' + r['estado']] for r in structure], [165,170,45,212,80]))
    def footer(canvas, doc):
        canvas.setFont('Helvetica',8); canvas.setFillColor(colors.HexColor('#40526A'))
        canvas.drawString(28,17,'PsiTwin Inventory - ' + label(data['generado_en']))
        canvas.drawRightString(814,17,f'Página {doc.page}')
    SimpleDocTemplate(out,pagesize=landscape(A4),leftMargin=28,rightMargin=28,topMargin=26,bottomMargin=34,title='PsiTwin Inventory - Resumen de inventario').build(story,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()


def generar_archivo(data, formato):
    if len(data['ubicaciones']) > 10000: raise ValueError('Acote el reporte a un almacén o producto (máximo 10000 ubicaciones)')
    return (excel(data), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'inventario.xlsx') if formato == 'xlsx' else (pdf(data), 'application/pdf', 'inventario.pdf')
