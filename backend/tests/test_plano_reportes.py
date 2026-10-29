import asyncio
from datetime import datetime, timezone, timedelta
from io import BytesIO
from unittest.mock import Mock
import smtplib
import pytest
from fastapi import HTTPException
from app.dto.auth_dto import UsuarioActual
from app.dto.inventario_consulta_dto import InventarioFiltros, ReporteSolicitud
from app.dto.rack_dto import RackCreate
from app.services.jerarquia_service import JerarquiaService
from app.services.rack_service import RackService
from app.services.inventario_consulta_service import consultar_inventario
from app.services.inventario_reportes_service import preparar_datos, InventarioReportesService
from app.services.inventario_reportes_archivos import generar_archivo
from app.services.smtp_transport import DeliveryUncertain
from app.config.settings import settings
from tests.test_google_auth import isolated_database
from tests.test_abastecimiento_integral import seed

async def sample(db):
    supplier, operator, company = await seed(db)
    await db.empresas.update_one({'id_empresa':'e'}, {'$set':{'nombre':'Empresa de prueba controlada','nombre_ceo':'Responsable de prueba','correo':'reportes@example.test','correo_verificado':True}})
    levels = await JerarquiaService.asegurar_niveles(await db.racks.find_one({'id_rack':'r'},{'_id':0}))
    await db.detalles_inventario.update_one({'id_detalle_inventario':'di'}, {'$set':{'id_nivel':levels[1]['id_nivel'],'stock_actual':10,'stock_minimo':15,'stock_maximo':20}})
    await db.detalles_inventario.insert_one({'id_detalle_inventario':'di2','id_inventario':'i','id_producto':'p','id_rack':'r','id_nivel':levels[2]['id_nivel'],'stock_actual':15,'stock_minimo':5,'stock_maximo':10,'fecha_actualizacion':datetime.now(timezone.utc)})
    await db.productos.insert_one({'id_producto':'u','id_empresa':'e','nombre':'Fresas de prueba','unidad':'unidad','familia':'Frutas','categoria':'Frutos rojos'})
    await db.detalles_inventario.insert_one({'id_detalle_inventario':'di3','id_inventario':'i','id_producto':'u','id_rack':'r','id_nivel':levels[0]['id_nivel'],'stock_actual':0,'stock_minimo':0,'stock_maximo':0})
    return supplier,operator,company,levels


def smtp_config(monkeypatch):
    for k,v in {'email_enabled':True,'smtp_user':'psitwin2026@gmail.com','smtp_app_password':'synthetic-test-only','smtp_host':'smtp.gmail.com','smtp_use_tls':True,'smtp_port':587}.items(): monkeypatch.setattr(settings,k,v)


def test_http_map_exports_permissions_and_expired_confirmation(monkeypatch):
    import httpx
    from fastapi import FastAPI
    from app.controllers import inventario_consulta_controller as controller
    from app.dependencies.dependencies import require_empresa, require_empresa_u_operario
    async def run():
        async with isolated_database(monkeypatch) as db:
            _, operator, company, levels = await sample(db)
            sender = Mock()
            monkeypatch.setattr(controller, 'reportes', InventarioReportesService(sender))
            instance = FastAPI(); instance.include_router(controller.router)
            instance.dependency_overrides[require_empresa] = lambda: company
            instance.dependency_overrides[require_empresa_u_operario] = lambda: company
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=instance), base_url='http://test') as client:
                plan = await client.get('/almacenes/w/plano')
                assert plan.status_code == 200 and len(plan.json()['niveles']) == 4
                page = await client.get('/inventario-consultas/resumen?limite=1&pagina=2')
                assert page.json()['total'] == 2 and len(page.json()['productos']) == 1
                for fmt, magic in [('pdf', b'%PDF'), ('xlsx', b'PK')]:
                    response = await client.get('/inventario-reportes/exportar', params={'formato':fmt,'id_nivel':levels[1]['id_nivel']})
                    assert response.status_code == 200 and response.content.startswith(magic)
                    assert response.headers['cache-control'] == 'no-store'
                assert (await client.get('/inventario-reportes/exportar?formato=exe')).status_code == 422
                draft = (await client.post('/inventario-reportes/solicitudes',json={'formato':'pdf','clave_idempotencia':'expiry-http-key'})).json()
                assert not sender.called
                await db.reportes_inventario.update_one({'id_reporte':draft['id_reporte']},{'$set':{'expira_en':datetime.now(timezone.utc)-timedelta(seconds=1)}})
                assert (await client.post(f"/inventario-reportes/envios/{draft['id_reporte']}/confirmar",json={'confirmar':True})).status_code == 410
                assert not sender.called
                instance.dependency_overrides[require_empresa_u_operario] = lambda: operator
                assert (await client.get('/almacenes/foreign/plano')).status_code in (403,404)
                instance.dependency_overrides[require_empresa] = lambda: operator
                assert (await client.get('/inventario-reportes/destinatario')).status_code == 403
    asyncio.run(run())


def test_technical_inventory_failure_is_not_an_empty_result(monkeypatch):
    from pymongo.errors import ServerSelectionTimeoutError
    from app.services import inventario_consulta_service as service
    def unavailable(): raise ServerSelectionTimeoutError('controlled failure without credentials')
    monkeypatch.setattr(service, 'get_database', unavailable)
    with pytest.raises(HTTPException) as error:
        asyncio.run(service.consultar_inventario(UsuarioActual(rol='EMPRESA',id_empresa='e')))
    assert error.value.status_code == 503


def test_saiki_names_real_products_paginates_and_honors_selected_level(monkeypatch):
    from app.services.saiki_service import SaikiService
    from app.dto.saiki_dto import SaikiAskRequest
    async def run():
        async with isolated_database(monkeypatch) as db:
            _, _, company, levels = await sample(db)
            service = SaikiService()
            async def forbidden(*args, **kwargs): pytest.fail('Inventory facts must not wait for or be invented by an LLM')
            monkeypatch.setattr(service,'_call_ollama_chat',forbidden)
            first = await service.ask(SaikiAskRequest(message='Saiki, ¿qué productos tengo?',page_size=1),company)
            second = await service.ask(SaikiAskRequest(message='Enumera productos de inventario, página 2',page_size=1),company)
            assert first['data']['paginacion']['total'] == 2
            assert len(first['data']['inventario_detallado']) == 1
            assert first['data']['inventario_detallado'][0]['id_producto'] != second['data']['inventario_detallado'][0]['id_producto']
            selected = await service.ask(SaikiAskRequest(message='¿Qué productos hay aquí?',context={'ubicacion':{'id_nivel':levels[1]['id_nivel']}}),company)
            assert 'Mango Kent: 10 kg' in selected['answer']
            assert selected['data']['filtros']['id_nivel'] == levels[1]['id_nivel']
            assert all(p['id_producto'] == 'p' for p in selected['data']['inventario_detallado'])
    asyncio.run(run())


def test_racks_get_exact_different_levels_and_zone_limits(monkeypatch):
    async def run():
        async with isolated_database(monkeypatch) as db:
            await seed(db)
            await db.zonas.update_one({'id':'z'},{'$set':{'max_racks':4,'capacidad':4}})
            racks = [await RackService().crear_rack(RackCreate(id_zona='z',codigo=f'R{i+2}',capacidad=100,max_niveles=n)) for i,n in enumerate([5,2,3])]
            for rack,n in zip(racks,[5,2,3]):
                levels=await JerarquiaService.asegurar_niveles(rack)
                assert [r['numero'] for r in levels] == list(range(1,n+1))
                assert len(await JerarquiaService.asegurar_niveles(rack)) == n
            with pytest.raises(HTTPException): await RackService().crear_rack(RackCreate(id_zona='z',codigo='R5',capacidad=100))
            assert await db.racks.count_documents({'id_zona':'z'}) == 4
    asyncio.run(run())


def test_consulta_consolidates_levels_units_thresholds_and_scope(monkeypatch):
    async def run():
        async with isolated_database(monkeypatch) as db:
            _,operator,company,levels=await sample(db)
            result=await consultar_inventario(company)
            assert result['resumen']['productos_distintos']==2
            assert result['resumen']['totales_por_unidad']=={'kg':25,'unidad':0}
            assert {a['tipo'] for a in result['alertas']} >= {'BAJO_MINIMO','SOBRE_MAXIMO','AGOTADO'}
            assert not any(a['tipo']=='BAJO_MINIMO' and a['id_producto']=='u' for a in result['alertas'])
            chosen=await consultar_inventario(operator,InventarioFiltros(id_nivel=levels[1]['id_nivel']))
            assert chosen['productos'][0]['stock_actual']==10
            assert len(chosen['ubicaciones'])==1
            await db.detalles_inventario.update_one({'id_detalle_inventario':'di'},{'$inc':{'stock_actual':-2}})
            assert (await consultar_inventario(company))['resumen']['totales_por_unidad']['kg']==23
            foreign=UsuarioActual(rol='EMPRESA',id_empresa='foreign')
            with pytest.raises(HTTPException): await consultar_inventario(foreign,InventarioFiltros(id_almacen='w'))
            with pytest.raises(HTTPException): await JerarquiaService.eliminar_nivel(company,levels[1]['id_nivel'])
    asyncio.run(run())


def test_excel_pdf_and_empty_inventory_are_valid_not_duplicated(monkeypatch):
    async def run():
        from openpyxl import load_workbook
        from pypdf import PdfReader
        async with isolated_database(monkeypatch) as db:
            _,_,company,_=await sample(db)
            data=await preparar_datos(company,InventarioFiltros())
            xlsx,_,_=generar_archivo(data,'xlsx'); pdf,_,_=generar_archivo(data,'pdf')
            wb=load_workbook(BytesIO(xlsx),data_only=False)
            assert wb.sheetnames==['Resumen','Por producto','Por ubicación','Alertas','Plano']
            assert wb['Por producto']['F8'].value==25
            assert wb['Por producto']['F8'].data_type=='n'
            assert wb['Por ubicación'].max_row==9
            assert 'Mango Kent' in ''.join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
            await db.productos.update_one({'id_producto':'p'},{'$set':{'nombre':'=HYPERLINK("https://invalid.test")'}})
            data=await preparar_datos(company,InventarioFiltros())
            wb=load_workbook(BytesIO(generar_archivo(data,'xlsx')[0]),data_only=False)
            assert all(c.data_type!='f' for ws in wb for row in ws for c in row)
            await db.detalles_inventario.delete_many({})
            data=await preparar_datos(company,InventarioFiltros())
            assert data['resumen']['productos_distintos']==0
            assert len(PdfReader(BytesIO(generar_archivo(data,'pdf')[0])).pages)>=1
    asyncio.run(run())


@pytest.mark.parametrize('format',['pdf','xlsx'])
def test_report_email_confirmation_is_single_and_registered(monkeypatch,format):
    async def run():
        async with isolated_database(monkeypatch) as db:
            _,operator,company,_=await sample(db); smtp_config(monkeypatch)
            sender=Mock(); service=InventarioReportesService(sender)
            request=ReporteSolicitud(formato=format,clave_idempotencia='test-mail-unique')
            action=await service.solicitar(company,request)
            assert action['estado']=='CONFIRMAR' and action['destinatario']=='reportes@example.test'
            sender.assert_not_called()
            duplicate=await service.solicitar(company,request)
            assert duplicate['id_reporte']==action['id_reporte']
            accepted=await service.confirmar(company,action['id_reporte'])
            assert accepted['estado']=='ACEPTADO_SMTP'
            await service.confirmar(company,action['id_reporte'])
            sender.assert_called_once()
            message=sender.call_args.args[0]
            assert message['To']=='reportes@example.test' and message['From']=='psitwin2026@gmail.com'
            attachment=list(message.iter_attachments())[0]
            assert attachment.get_filename()=='inventario.'+format and attachment.get_payload(decode=True)
            with pytest.raises(HTTPException): await service.solicitar(operator,request)
            with pytest.raises(HTTPException): await service.obtener(UsuarioActual(rol='EMPRESA',id_empresa='foreign'),action['id_reporte'])
    asyncio.run(run())


@pytest.mark.parametrize('error,state,code',[(smtplib.SMTPAuthenticationError(535,b'synthetic'), 'ERROR','SMTP_AUTENTICACION'), (ConnectionError(), 'ERROR','SMTP_CONEXION'), (DeliveryUncertain(), 'REVISAR_ENVIO','CONFIRMACION_SMTP_DESCONOCIDA')])
def test_mail_errors_do_not_claim_success_or_retry(monkeypatch,error,state,code):
    async def run():
        async with isolated_database(monkeypatch) as db:
            _,_,company,_=await sample(db); smtp_config(monkeypatch)
            sender=Mock(side_effect=error); service=InventarioReportesService(sender)
            action=await service.solicitar(company,ReporteSolicitud(formato='pdf',clave_idempotencia='test-error-key'))
            result=await service.confirmar(company,action['id_reporte'])
            assert result['estado']==state and result['codigo_error']==code
            await service.confirmar(company,action['id_reporte'])
            sender.assert_called_once()
    asyncio.run(run())


def test_saiki_three_turns_authorized_confirmation_and_cancellation(monkeypatch):
    async def run():
        from app.services import saiki_reportes
        from app.services.saiki_service import SaikiService
        from app.dto.saiki_dto import SaikiAskRequest
        async with isolated_database(monkeypatch) as db:
            _,_,company,_=await sample(db); smtp_config(monkeypatch)
            sender=Mock(); monkeypatch.setattr(saiki_reportes,'service',InventarioReportesService(sender))
            service=SaikiService()
            first=await service.ask(SaikiAskRequest(message='Saiki, mándame el resumen de inventario a mi correo'),company)
            action=first['data']['reporte_accion']; assert action['estado']=='ELEGIR_FORMATO'
            second=await service.ask(SaikiAskRequest(message='PDF',context={'reporte_accion':action['id_reporte']}),company)
            assert second['data']['reporte_accion']['estado']=='CONFIRMAR'; sender.assert_not_called()
            third=await service.ask(SaikiAskRequest(message='Sí',context={'reporte_accion':action['id_reporte']}),company)
            assert third['data']['reporte_accion']['estado']=='ACEPTADO_SMTP'; sender.assert_called_once()
            await service.ask(SaikiAskRequest(message='Sí',context={'reporte_accion':action['id_reporte']}),company)
            sender.assert_called_once()
            bad=await service.ask(SaikiAskRequest(message='Envíame el reporte a intruso@example.test'),company)
            assert 'destinatarios' in bad['answer']; sender.assert_called_once()
            fresh=await service.ask(SaikiAskRequest(message='Mándame el reporte PDF al correo'),company)
            cancel=await service.ask(SaikiAskRequest(message='Cancelar',context={'reporte_accion':fresh['data']['reporte_accion']['id_reporte']}),company)
            assert cancel['data']['reporte_accion']['estado']=='CANCELADO'
    asyncio.run(run())
